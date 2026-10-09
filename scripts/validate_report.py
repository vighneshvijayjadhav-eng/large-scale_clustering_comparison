"""Check report structure and evidence without fitting any estimator."""

import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"


def main():
    source = (REPORTS / "source/report.md").read_text(encoding="utf-8")
    sections = source.split("@@ ")[1:]
    reader = PdfReader(REPORTS / "DWM_Mini_Project_Report.pdf")
    assert len(reader.pages) == len(sections) == 34
    for index, (section, page) in enumerate(zip(sections, reader.pages, strict=True)):
        text = page.extract_text()
        assert len(text) > 250, f"Unexpected blank or overflow page {index + 1}"
        if index:
            assert section.splitlines()[0] in text, f"Page boundary mismatch {index + 1}"
        assert abs(float(page.mediabox.width) - 595.276) < 0.01
    document = Document(REPORTS / "DWM_Mini_Project_Report.docx")
    assert len(document.tables) == 11
    assert len(document.inline_shapes) == 23  # 17 numbered figures + six equations
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert document.styles["Normal"].font.size.pt == 12
    assert document.styles["Normal"].paragraph_format.line_spacing == 1.5
    assert "TOC" in document._element.xml
    assert "PAGE" in document.sections[0].footer._element.xml
    assert sum(p.style.name == "Title" for p in document.paragraphs) == 1
    abstract = next(s for s in sections if s.startswith("Abstract\n"))
    abstract_words = len(abstract.split("Keywords:")[0].split()) - 1
    assert 250 <= abstract_words <= 350
    with (REPORTS / "benchmarks.csv").open(encoding="utf-8-sig") as stream:
        benchmarks = list(csv.DictReader(stream))
    for size in [50000, 100000, 250000, 500000, 1000000]:
        assert len([r for r in benchmarks if int(r["rows"]) == size and r["status"] == "ok"]) == 2
    assert any(
        r["algorithm"] == "KMeans" and r["status"] == "skipped_memory_preflight" for r in benchmarks
    )
    files = [
        "benchmarks.csv",
        "batch_study.csv",
        "profiles.csv",
        "k_selection.csv",
        "selection.json",
        "provenance.json",
        "model_metadata.json",
        "runtime.json",
        "inference_smoke.json",
        "source/report.md",
        "DWM_Mini_Project_Report.docx",
        "DWM_Mini_Project_Report.pdf",
    ]
    checks = {
        "pdf_pages": len(reader.pages),
        "docx_pagination": "Not rendered: LibreOffice unavailable; 34 explicit source pages",
        "numbered_figures": len(re.findall(r"\{\{figure\|", source)),
        "numbered_tables": len(document.tables),
        "equations": 6,
        "abstract_words": abstract_words,
        "sha256": {
            name: hashlib.sha256((REPORTS / name).read_bytes()).hexdigest() for name in files
        },
    }
    (REPORTS / "source/validation_checks.json").write_text(
        json.dumps(checks, indent=2), encoding="utf-8"
    )
    print(json.dumps({k: v for k, v in checks.items() if k != "sha256"}, indent=2))


if __name__ == "__main__":
    main()
