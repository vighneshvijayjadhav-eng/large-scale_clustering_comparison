"""Build editable DOCX and PDF from one evidence-backed report source."""

import csv
import json
import re
import shutil
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as PDFImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIGURES = REPORTS / "figures"


def rows(name):
    with (REPORTS / name).open(encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def table_data(key):
    benchmark = rows("benchmarks.csv")
    paired = [
        r
        for r in benchmark
        if r["status"] == "ok" and r["streaming"] == "false" and int(r["rows"]) <= 1000000
    ]
    if key == "counts":
        return [
            ["Population / exclusion", "Records"],
            ["Raw January source", "3,475,226"],
            ["Accepted", "3,241,580"],
            ["Rejected total", "233,646"],
            ["Invalid charges", "127,550"],
            ["Invalid distance", "64,474"],
            ["Invalid duration", "41,488"],
            ["Implausible speed", "134"],
        ]
    if key == "statistics":
        return [["Statistic", "Miles", "Minutes", "mph", "USD/mile"]] + [
            [r["statistic"]]
            + [
                f"{float(r[c]):,.2f}"
                for c in ["trip_distance", "duration_minutes", "speed_mph", "fare_per_mile"]
            ]
            for r in rows("feature_statistics.csv")
            if r["statistic"] in ["mean", "std", "min", "50%", "max", "median"]
        ]
    if key == "selection":
        return [["K", "Inertia / row", "Silhouette", "Stability ARI", "Min. share", "Eligible"]] + [
            [
                r["k"],
                f"{float(r['normalized_inertia']):.4f}",
                f"{float(r['silhouette']):.4f}",
                f"{float(r['stability_ari']):.4f}",
                f"{100 * float(r['min_cluster_share']):.3f}%",
                r["eligible"],
            ]
            for r in rows("k_selection.csv")
        ]
    if key == "runtime":
        return [
            ["Setting", "Recorded value"],
            ["Operating system", "Windows 11, build 26200"],
            ["Memory / logical CPUs", "5.889 GiB / 8; CPU model not recorded"],
            ["Python / numerical threads", "3.12.13 / 2"],
            ["NumPy / Polars", "2.5.3 / 1.44.2"],
            ["scikit-learn / UMAP", "1.9.1 / 0.5.12"],
            ["K / seed / features", "3 / 42 / 6 standardized float32 columns"],
            ["Initialization / n_init", "k-means++ / 10"],
            ["KMeans / MiniBatch max_iter", "300 / 300"],
            ["Main MiniBatch size", "1,000"],
            ["Silhouette evaluation", "2,000 reproducibly selected records"],
        ]
    if key in ["performance", "quality"]:
        header = (
            ["Rows", "Algorithm", "Fit seconds", "Peak RSS MiB"]
            if key == "performance"
            else ["Rows", "Algorithm", "Inertia / row", "Silhouette", "ARI vs KM"]
        )
        return [header] + [
            [f"{int(r['rows']) // 1000}K", "KM" if r["algorithm"] == "KMeans" else "MB"]
            + (
                [f"{float(r['fit_seconds']):.4f}", f"{float(r['peak_rss_mb']):.2f}"]
                if key == "performance"
                else [
                    f"{float(r['normalized_inertia']):.4f}",
                    f"{float(r['silhouette']):.4f}",
                    f"{float(r['ari_vs_kmeans']):.4f}" if r["ari_vs_kmeans"] else "Reference",
                ]
            )
            for r in paired
        ]
    if key == "batch":
        return [["Batch size", "Fit seconds", "Inertia / row", "Silhouette", "ARI vs KM"]] + [
            [r["batch_size"]]
            + [
                f"{float(r[c]):.4f}"
                for c in ["fit_seconds", "normalized_inertia", "silhouette", "ari_vs_kmeans"]
            ]
            for r in rows("batch_study.csv")
        ]
    if key == "profiles":
        return [["Model / ID", "Trips", "Share %", "Miles", "Minutes", "mph", "USD/mi"]] + [
            [
                ("KM" if r["algorithm"] == "KMeans" else "MB") + " / " + r["cluster"],
                f"{int(r['size']):,}",
                f"{int(r['size']) / 1000:.2f}",
            ]
            + [
                f"{float(r[c]):.2f}"
                for c in ["trip_distance", "duration_minutes", "speed_mph", "fare_per_mile"]
            ]
            for r in rows("profiles.csv")
        ]
    if key == "anomalies":
        return [
            ["Model", "Cluster 0", "Cluster 1", "Cluster 2", "Total / training"],
            ["KMeans", "233", "34", "733", "1,000 / 100,000"],
            ["MiniBatch", "18", "248", "734", "1,000 / 100,000"],
        ]
    if key == "directory":
        return [
            ["Path", "Purpose"],
            ["src/taxi/", "Validation, fitting, inference, benchmark and plotting"],
            ["app.py; .streamlit/", "Local dashboard and theme"],
            ["config/; tests/", "Reproducible settings and offline tests"],
            ["reports/", "Measured CSV/JSON results and this report"],
            ["scripts/; notebooks/", "Smoke checks, report generation and Colab fallback"],
            ["data/; artifacts/", "Ignored local data and saved model binaries"],
        ]
    if key == "holdout":
        sample = json.loads((REPORTS / "source/holdout_example.json").read_text(encoding="utf-8"))
        return [["row_id", "Miles", "Minutes", "Hour", "Cluster", "Anomaly ratio"]] + [
            [
                r["row_id"],
                r["trip_distance"],
                f"{float(r['duration_minutes']):.3f}",
                r["pickup_hour"],
                r["cluster"],
                f"{float(r['anomaly_ratio']):.4f}",
            ]
            for r in sample
        ]
    raise ValueError(key)


def prepare_captures():
    for source, target in [
        ("overview.jpg", "ui_overview.jpg"),
        ("cluster_selection.jpg", "ui_comparison.jpg"),
    ]:
        destination = FIGURES / target
        if not destination.exists():
            shutil.copyfile(ROOT / "docs/screenshots" / source, destination)
    missing = []
    for name, title in [
        ("ui_anomalies.jpg", "Anomaly Explorer"),
        ("ui_predict.jpg", "Predict New Trips"),
        ("ui_methodology.jpg", "Methodology / Experiment Details"),
    ]:
        if not (FIGURES / name).exists():
            image = Image.new("RGB", (1200, 540), "#f1f5f9")
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 32)
            draw.multiline_text(
                (70, 150),
                f"SCREENSHOT PLACEHOLDER\n{title}\n\nA verified page capture was unavailable for this report.\nThis panel is not a reconstructed application screenshot.",
                fill="#334155",
                font=font,
                spacing=14,
            )
            image.save(FIGURES / name, quality=95)
            missing.append(name)
    return missing


def field(paragraph, instruction):
    run = paragraph.add_run()
    item = OxmlElement("w:fldSimple")
    item.set(qn("w:instr"), instruction)
    run._r.addnext(item)


def main():
    prepare_captures()
    text = (REPORTS / "source/report.md").read_text(encoding="utf-8")
    pages = [
        (p.split("\n", 1)[0].strip(), p.split("\n", 1)[1].strip()) for p in text.split("@@ ")[1:]
    ]
    figures, tables = [], []
    for page, (_, body) in enumerate(pages, 1):
        for match in re.finditer(r"\{\{(figure|table)\|([^}]+)\}\}", body):
            parts = match[2].split("|")
            (figures if match[1] == "figure" else tables).append((parts[1], page))
    for family, filename in [
        ("TimesNewRoman", "times.ttf"),
        ("TimesNewRoman-Bold", "timesbd.ttf"),
        ("TimesNewRoman-Italic", "timesi.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(family, "C:/Windows/Fonts/" + filename))
    pdfmetrics.registerFontFamily(
        "TimesNewRoman",
        normal="TimesNewRoman",
        bold="TimesNewRoman-Bold",
        italic="TimesNewRoman-Italic",
        boldItalic="TimesNewRoman-Bold",
    )
    body_style = ParagraphStyle(
        "Body",
        fontName="TimesNewRoman",
        fontSize=12,
        leading=18,
        alignment=TA_JUSTIFY,
        spaceAfter=7,
    )
    heading_style = ParagraphStyle(
        "Heading",
        parent=body_style,
        fontName="TimesNewRoman-Bold",
        fontSize=16,
        leading=20,
        spaceAfter=12,
        alignment=0,
    )
    sub_style = ParagraphStyle(
        "Subheading", parent=heading_style, fontSize=12, leading=15, spaceBefore=4, spaceAfter=5
    )
    caption_style = ParagraphStyle(
        "Caption", parent=body_style, fontSize=10, leading=12, alignment=TA_CENTER, spaceAfter=8
    )
    cell_style = ParagraphStyle(
        "Cell", parent=body_style, fontSize=9.5, leading=11.5, spaceAfter=0, alignment=0
    )
    compact_style = ParagraphStyle(
        "Compact", parent=body_style, fontSize=11, leading=14, spaceAfter=5, alignment=0
    )
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(size)
        style.font.color.rgb = __import__("docx").shared.RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_after = Pt(12 if name == "Heading 1" else 5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, "PAGE")
    story = []
    fi = ti = 0
    width = A4[0] - 115.2
    for index, (title, body) in enumerate(pages):
        if index:
            doc.add_page_break()
            story.append(PageBreak())
        if index == 0:
            story.append(Spacer(1, 45))
        else:
            doc.add_heading(title, level=1)
            story.append(Paragraph(escape(title), heading_style))
        body = re.sub(r"(## [^\n]+)\n(?!\n)", lambda match: match[1] + "\n\n", body)
        for block in body.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("{{"):
                parts = block[2:-2].split("|")
                kind = parts[0]
                if kind in ["toc", "figlist", "tablelist"]:
                    entries = (
                        [(t, i + 1) for i, (t, _) in enumerate(pages) if i >= 7]
                        if kind == "toc"
                        else [
                            (
                                f"{'Figure' if kind == 'figlist' else 'Table'} {i + 1}. {caption}",
                                page,
                            )
                            for i, (caption, page) in enumerate(
                                figures if kind == "figlist" else tables
                            )
                        ]
                    )
                    if kind == "toc":
                        toc_start = doc.add_paragraph().add_run()._r
                        begin = OxmlElement("w:fldChar")
                        begin.set(qn("w:fldCharType"), "begin")
                        instruction = OxmlElement("w:instrText")
                        instruction.text = 'TOC \\o "1-1" \\h \\z \\u'
                        separate = OxmlElement("w:fldChar")
                        separate.set(qn("w:fldCharType"), "separate")
                        for element in [begin, instruction, separate]:
                            toc_start.append(element)
                    for label, page in entries:
                        line = f"{label}  ·  {page}"
                        p = doc.add_paragraph(line)
                        p.paragraph_format.line_spacing = 1
                        p.paragraph_format.space_after = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(10)
                        story.append(Paragraph(escape(line), compact_style))
                    if kind == "toc":
                        end = OxmlElement("w:fldChar")
                        end.set(qn("w:fldCharType"), "end")
                        doc.add_paragraph().add_run()._r.append(end)
                elif kind == "figure":
                    fi += 1
                    path = FIGURES / parts[1]
                    with Image.open(path) as im:
                        w, h = im.size
                    height = min(float(parts[3]) * 72, width * h / w)
                    image_width = height * w / h
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1
                    p.paragraph_format.space_after = Pt(2)
                    p.add_run().add_picture(str(path), width=Inches(image_width / 72))
                    caption = f"Figure {fi}. {parts[2]}."
                    p = doc.add_paragraph(caption, "Caption")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1
                    story.extend(
                        [
                            PDFImage(str(path), width=image_width, height=height),
                            Paragraph(escape(caption), caption_style),
                        ]
                    )
                elif kind == "table":
                    ti += 1
                    caption = f"Table {ti}. {parts[2]}."
                    doc.add_paragraph(caption, "Caption")
                    data = table_data(parts[1])
                    table = doc.add_table(rows=0, cols=len(data[0]))
                    table.style = "Light Shading Accent 1"
                    for row in data:
                        cells = table.add_row().cells
                        for cell, value in zip(cells, row, strict=True):
                            cell.text = str(value)
                            for p in cell.paragraphs:
                                p.paragraph_format.line_spacing = 1
                                p.paragraph_format.space_after = Pt(3)
                                for run in p.runs:
                                    run.font.size = Pt(9.5)
                    pdf_table = Table(
                        [[Paragraph(escape(str(c)), cell_style) for c in row] for row in data],
                        colWidths=[width / len(data[0])] * len(data[0]),
                        repeatRows=1,
                    )
                    pdf_table.setStyle(
                        TableStyle(
                            [
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#94a3b8")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("TOPPADDING", (0, 0), (-1, -1), 4),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                            ]
                        )
                    )
                    story.extend(
                        [Paragraph(escape(caption), caption_style), pdf_table, Spacer(1, 7)]
                    )
                elif kind == "equation":
                    equation_path = FIGURES / f"equation_{parts[-1]}.png"
                    with Image.open(equation_path) as im:
                        equation_width = min(400, im.width * 72 / 240)
                        equation_height = equation_width * im.height / im.width
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1
                    p.add_run().add_picture(str(equation_path), width=Inches(equation_width / 72))
                    p.add_run(f"  ({parts[-1]})")
                    story.append(
                        Table(
                            [
                                [
                                    PDFImage(
                                        str(equation_path),
                                        width=equation_width,
                                        height=equation_height,
                                    ),
                                    Paragraph(f"({parts[-1]})", caption_style),
                                ]
                            ],
                            colWidths=[415, 40],
                        )
                    )
                elif kind == "code":
                    for line in parts[1].split("~"):
                        p = doc.add_paragraph(line)
                        p.paragraph_format.line_spacing = 1
                        p.paragraph_format.space_after = Pt(2)
                        for run in p.runs:
                            run.font.name, run.font.size = "Consolas", Pt(9)
                        story.append(Paragraph(escape(line), compact_style))
            elif block.startswith("## "):
                doc.add_heading(block[3:], level=2)
                story.append(Paragraph(escape(block[3:]), sub_style))
            else:
                value = block.replace("\n", "\n" if index == 0 else " ")
                p = doc.add_paragraph(value)
                style = body_style
                if index == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if value.startswith("Scalable"):
                        p.style = doc.styles["Title"]
                    style = ParagraphStyle(
                        "TitleBlock",
                        parent=body_style,
                        alignment=TA_CENTER,
                        spaceAfter=20,
                        fontSize=18 if value.startswith("Scalable") else 12,
                        leading=24 if value.startswith("Scalable") else 18,
                    )
                if title == "References":
                    p.paragraph_format.line_spacing = 1
                    style = compact_style
                story.append(Paragraph(escape(value).replace("\n", "<br/>"), style))
    doc.core_properties.title = (
        "Scalable Taxi Trip Pattern Mining Using K-Means and MiniBatch K-Means"
    )
    doc.core_properties.author = "Vighnesh Vijay Jadhav; Shreesh Jugade"
    doc.save(REPORTS / "DWM_Mini_Project_Report.docx")

    def page_footer(canvas, document):
        canvas.setFont("TimesNewRoman", 10)
        canvas.drawCentredString(A4[0] / 2, 27, str(document.page))

    pdf = SimpleDocTemplate(
        str(REPORTS / "DWM_Mini_Project_Report.pdf"),
        pagesize=A4,
        leftMargin=57.6,
        rightMargin=57.6,
        topMargin=50.4,
        bottomMargin=50.4,
        title=doc.core_properties.title,
        author=doc.core_properties.author,
    )
    pdf.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    (REPORTS / "source/build_manifest.json").write_text(
        json.dumps(
            {
                "source_pages": len(pages),
                "figures": len(figures),
                "tables": len(tables),
                "words": len(text.split()),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Built {len(pages)} source pages, {len(figures)} figures, {len(tables)} tables")


if __name__ == "__main__":
    main()
